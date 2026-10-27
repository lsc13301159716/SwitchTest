import json
import statistics
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
DEST = ROOT / '报告'
DEST.mkdir(exist_ok=True)
data = json.loads((ROOT/'out/analysis_20261014/analysis.json').read_text(encoding='utf-8'))
rounds, events = data['rounds'], data['events']
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
sec.left_margin, sec.right_margin = Cm(2.0), Cm(2.0)
sec.footer_distance = Cm(.8)
for name, size in [('Normal',10.5),('Title',22),('Subtitle',11),('Heading 1',15),('Heading 2',11.5)]:
    st = doc.styles[name]
    st.font.name = 'Microsoft YaHei'
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0,0,0)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.line_spacing = 1.15
    st.font.italic = False
    for border in list(st.element.xpath('./w:pPr/w:pBdr')):
        border.getparent().remove(border)
    snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'),'0')
    st.element.get_or_add_pPr().append(snap)
    if name.startswith('Heading'):
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.keep_with_next = True
foot = sec.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
foot.add_run('第 ')
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); foot._p.append(fld)
foot.add_run(' 页')
for r in foot.runs:
    r.font.size = Pt(9)

md=[]
def p(text, style=None):
    doc.add_paragraph(text, style)
    md.append(text+'\n')
def h(text, level=1):
    doc.add_heading(text,level)
    md.append('#'*(level+1)+' '+text+'\n')
def page():
    pass
def table(headers, rows, widths, size=9):
    t=doc.add_table(rows=1, cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for col,w in zip(t.columns,widths): col.width=Cm(w)
    for cell,w,text in zip(t.rows[0].cells,widths,headers):
        cell.width=Cm(w); cell.text=str(text)
    for vals in rows:
        cells=t.add_row().cells
        for c,w,val in zip(cells,widths,vals): c.width=Cm(w); c.text=str(val)
    for i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');pr.append(cant)
        if i==0:
            rep=OxmlElement('w:tblHeader');pr.append(rep)
        for c in row.cells:
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            cp.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');margins.append(el)
            cp.append(margins)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EDF2');cp.append(sh)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(0)
                para.paragraph_format.line_spacing=1.10
                snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');para._p.get_or_add_pPr().append(snap)
                para.alignment=WD_ALIGN_PARAGRAPH.CENTER if len(c.text)<28 else WD_ALIGN_PARAGRAPH.LEFT
                for run in para.runs:
                    run.font.size=Pt(size);run.bold=i==0
    md.append('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |')
    md.extend('| '+' | '.join(str(x).replace('\n','；') for x in row)+' |' for row in rows)
    md.append('')
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

doc.add_paragraph('WAGO与Weidmuller交换机\n转发可靠性对比试验报告','Title')
md.append('# WAGO与Weidmuller交换机转发可靠性对比试验报告\n')
p('版本 V0.1    编制日期 2026年10月14日','Subtitle')
p('试验日期：2026年10月13日至14日及后续长测日期【待补】\n试验地点：【待补】    编制人：【待补】    审核人：【待补】')
h('阶段性摘要')
p('已完成直连三种帧长各60秒测试，以及WAGO、Weidmuller两台交换机各16组端口接线、三种帧长的60秒测试，共99轮。直连未观察到超时事件；WAGO和Weidmuller各有4轮出现超时回包，分别涉及4个和5个迟到包。全部测试尝试最终取得回包，未记录本地发送失败或socket错误。')
p('8个异常轮次均在迟到附近出现超过50ms的实际发包间隙，且间隙与迟到RTT接近，提示测量端程序或系统暂停可能参与异常。当前证据不能确定异常责任端口或设备。重点端口长测及异常复核正在进行，最终使用或替换意见待结果补齐后形成。')
h('1 试验目的与范围')
p('本试验检查两台交换机各端口在指定UDP回显负载下的转发与回包情况，比较同条件下的超时事件、连续探测失败和端到端时延，并通过直连与重复测量复核异常，为设备使用及可能的替换评估提供试验依据。原装设备、拟替换设备及实际生产用途的对应关系为【待补】。')
p('当前范围为两台Windows PC、单流UDP、64/512/1518B请求帧长、目标1000pps的全口短测及后续600/1800秒重点测试。本阶段不覆盖16口同时满载、全部速率下的线速性能、30天老化或所有生产工况。高负荷单向flood测试如后续开展，应单独补充方法与结果，不与本报告echo超时统计混合。')

page();h('2 被测设备与测试环境')
h('2 1 设备信息',2)
table(['项目','已记录内容','需补充内容'],[
['WAGO交换机','设备标识 WAGO','型号、编号、固件、端口标称速率、配置'],
['Weidmuller交换机','设备标识 Weidmuller','型号、编号、固件、端口标称速率、配置'],
['PC A发送端','有线IP 192.168.214.230','型号、Windows与Python版本、网卡型号'],
['PC B接收端','有线IP 192.168.214.231','型号、Windows与Python版本、网卡型号'],
['网卡与线缆','两端有线连接','协商速率、USB扩展坞、线缆编号及长度'],
['测量软件','SwitchTest.py与switch_core.py','两端文件SHA256及部署版本'],
['现场条件','【待补】','供电、电源管理、后台负载、现场操作记录']],[3.2,5.3,8.5])
h('2 2 拓扑与端口定义',2)
p('直连路径为PC A网口连接PC B网口；经交换机路径为PC A连接A侧端口，PC B连接B侧端口。A、B是收发电脑角色，不表示交换机方向。端口对A13—B01表示PC A连接13号口、PC B连接1号口。直连文件中的A01_B01仅为命名标签，不解释为交换机端口。')
p('全口短测先将PC B固定在1号口，PC A依次接2至16号口；随后PC A接1号口、PC B接2号口。端口覆盖依据文件标签整理，实际接线由执行人员核对确认【待补】。')
h('2 3 可比性记录',2)
p('当前日志表明发送端主要参数相同，接收端两份日志均为普通server模式。WAGO于10月13日测量，Weidmuller于10月14日测量；两批并非交替测试。实际协商速率、版本哈希及电脑负载未随本批日志记录，需补录，并在后续配对复测中尽量保持相同。')

page();h('3 测试方法与判读规则')
h('3 1 UDP回显测量方法',2)
p('PC A发送包含会话标识及包号的UDP请求，PC B返回回显。程序以发送前单调时钟和发送端应用读取有效回包的时刻计算RTT，并按固定timeout判定是否按时收到。RTT包含两端软件、系统调度、协议栈、网卡与交换机路径的影响，不能单独解释为交换机内部转发时延。')
table(['参数','本批统一设置'],[['请求帧长','64、512、1518B'],['发包间隔与进程数','0.001s，单流 jobs=1，目标1000pps'],['正式发包时长','全口及现有直连60s；长测计划600/1800s'],['超时与收尾观察','timeout 0.05s；drain 0.5s'],['预热','warmup 3s上限，成功后可提前结束'],['接收端','192.168.214.231 UDP 5001，普通server'],['发送端源地址','192.168.214.230'],['原始数据','逐包CSV、迟到附件、发送端及服务端日志']],[5,12])
p('帧长按普通无VLAN的IPv4 UDP以太网请求帧计算，三种帧长分别对应18、466、1472B UDP载荷。回显返回报文头，回程不是同长度满帧。本测试为低负载回显测试，目标1000pps不代表100Mbps或1Gbps线速。')
h('3 2 事件与时延统计',2)
p('on_time表示在50ms门槛内读取回包；late表示超时后仍收到回包；unreturned表示观察结束仍未收到；send_failed表示本地提交失败。迟到仍计入超时事件。连续至少2包超时记为连续探测失败段，不自动解释为物理链路中断。')
p('按时RTT计算p50、p90、p99，采用排序后位置(n−1)×p的线性插值。原始最大收到RTT包含迟到包。裁尾最大值移除最大ceil(n×0.001)个已收到RTT后取最大值，仅辅助比较。按时分位数受50ms门槛截断；实际pps采用尝试数除以正式发包时长。')
h('3 3 暂定判读规则',2)
p('暂按现有判据建议：有效完成且零失败为PASS；全为孤立事件、无本地发送失败且失败比例不超过0.001%为REVIEW；连续至少2包失败或比例超过0.001%为FAIL；明细不完整、测试中止或关键条件不可比时按具体用途列为INVALID。该规则待验收方确认【待补】，FAIL表示本轮不达标，不表示设备硬件故障或直接拒收。50ms也尚未确定为实际业务允许停顿。')

page();h('4 测试方案与执行记录')
table(['阶段','接线范围','帧长B','每帧时长','执行状态'],[
['直连短基线','PC A直连PC B','64/512/1518','60s','已完成3轮'],
['全口短测','每台16组端口接线','64/512/1518','60s','各完成48轮'],
['直连长基线','PC A直连PC B','64/512/1518','600s','结果待补'],
['异常复核','第6章异常端口对及对照','优先异常帧长','600s','正在测试 结果待补'],
['代表性长测','1—8、1—16、8—16','64/512/1518','600s','结果待补'],
['延长验证','重复异常及关键生产组合','按用例确定','1800s','按结果选择 待补']],[2.7,5.0,2.8,2.5,4.0])
h('4 1 已完成测试清单',2)
p('直连3轮正式发包时间均为60.000s。两台交换机均完成A02—B01至A16—B01及A01—B02的三种帧长测试，各48轮，每轮60.000s。合计99轮，纯正式发包时间99分钟，不含换线、预热、服务等待及后续长测。各口详细结果见附录A、B。')
h('4 2 长测与复核执行原则',2)
p('后续先建立当日600秒直连基线，再按异常优先级进行原接线600秒复测，并在另一台设备上用相同端口对、帧长及参数交替对照。存在差异时补至3组配对复测，记录轮次、开始时间及现场动作。3组仅作为初步复核，不宣称统计显著。')
p('原接线异常复现时，改变参考口或逐项更换已编号线缆进行定位。每次只改变一个条件并保留全部原始记录。多个异常共享B01，需同时考虑参考口、线缆和端点因素。一次复测无异常不撤销首次事件。')
p('代表性长测保留1—8、1—16、8—16组合；原规程中1—15为历史异常候选，可按实际生产重要性决定保留或后排。本批15—1的三种短测均无失败。生产关键端口不在这些组合时应追加【待补生产口清单】。')
h('4 3 记录与归档要求',2)
p('每轮保存唯一编号的CSV、迟到附件、日志、端口接线和完整参数。命名差异通过原始文件索引关联，不重命名或覆盖既有证据。报告补充时按实际完成状态更新，进行中轮次不得填零或写为通过。')

page();h('5 测试结果')
h('5 1 短测汇总',2)
rows=[]
for dev in ('DIRECT','WAGO','Weidmuller'):
    rs=[r for r in rounds if r['device']==dev]
    rows.append([dev,len(rs),f"{sum(r['n'] for r in rs):,}",sum(r['late']>0 for r in rs),sum(r['late'] for r in rs),sum(r['unreturned'] for r in rs),max(r['max_run'] for r in rs)])
table(['设备','轮数','尝试总数','异常轮','迟到包','未回包','最长连续'],rows,[3.3,1.3,3.3,2.1,2.0,2.0,3.0])
p('99份发送端日志与CSV的尝试数、按时回包数和失败数一致。CSV worker均为0，包号连续且发送时刻单调，迟到附件数量匹配。两台各44轮没有超时，另各4轮出现迟到；全部轮次本地发送失败和socket错误均为0。')
p('按第3章暂定规则，91轮为PASS，8轮为FAIL。每个单包迟到轮次约6万次尝试，超时比例约0.0017%，超过0.001%暂定界限；Weidmuller A13—B01的1518B还出现连续2包超时。该标签仅评价本轮观测，事件原因仍待复核。不得用整批合并比例稀释异常轮次。')
h('5 2 主体与尾部时延',2)
lat=[]
for dev in ('DIRECT','WAGO','Weidmuller'):
    for frame in (64,512,1518):
        rs=[r for r in rounds if r['device']==dev and r['frame']==frame]
        lat.append([dev,frame]+[f"{statistics.median(r[k] for r in rs):.3f}" for k in ('p50_ms','p90_ms','p99_ms')])
table(['设备','帧长B','p50 ms','p90 ms','p99 ms'],lat,[4.4,2.5,3.3,3.3,3.5])
p('表中为同帧长各轮分位数的中位数，直连每帧长仅1轮。Weidmuller p50较WAGO高约2%～5%，p90高约1%～3%，p99高约24%～70%。两台分日期测试，不能把该差异直接解释为设备固有性能差异。')

page();h('5 3 异常轮次明细',2)
exc=[]
for r in rounds:
    if r['late']:
        es=[e for e in events if e['file']==r['file']]
        exc.append([r['device'],f"{r['a']}—{r['b']}",r['frame'],', '.join(str(e['seq']) for e in es),'/'.join(f"{e['late_rtt_ms']:.3f}" for e in es),f"{max(e['nearby_send_gap_ms'] for e in es):.3f}"])
table(['设备','A—B口','帧长B','迟到包号','迟到RTT ms','附近间隙ms'],exc,[3.3,1.7,1.7,3.5,3.6,3.2],8.5)
p('以上8个异常轮次共9个迟到包。除Weidmuller A13—B01的1518B为连续2包外，其余均为单包迟到。附近间隙是根据CSV实际send_s差值计算，不使用包号乘以设定间隔估计。逐包发送时刻和原始文件名见exception_packets.csv。')
h('5 4 长测结果填写区',2)
p('本表供已完成长测及复核轮次补录。尚未填入的结果表示待补，不表示零事件。帧长、时长不同的轮次分别记录。')
table(['设备及端口对','帧长B 时长s','尝试数 迟到 未回包','最长连续与RTT','状态及轮次'],[
['直连','【待补】','【待补】','【待补】','【待补】'],
['Weidmuller 13—1','【待补】','【待补】','【待补】','【待补】'],
['两台 2—1','【待补】','【待补】','【待补】','【待补】'],
['其他异常端口对','【待补】','【待补】','【待补】','【待补】'],
['代表性与生产关键口','【待补】','【待补】','【待补】','【待补】']],[4.0,2.8,3.8,3.5,2.9])
p('各轮补录后，同时更新实际pps、按时p50/p90/p99、最大收到RTT、现场动作和相应原始文件索引。')

page();h('6 异常分析与复核')
h('6 1 迟到与发送间隙的关联',2)
p('8个异常轮次均出现超过50ms的实际发送间隙，且间隙与迟到RTT接近。WAGO A11—B01的64B迟到RTT为126.759ms，迟到包到下一次发送间隙为126.532ms；A16—B01的512B相应为124.221ms与123.920ms。')
p('Weidmuller A13—B01的1518B中，16366与16367号包约隔1ms发出，随后到16368号包出现116.412ms间隙；两包RTT为117.532ms和116.544ms。全批次最大实际发包间隙超过50ms的轮次恰好为这8个异常轮次。')
p('发送间隙表明发送端在该时段未按1ms持续发包，提示系统调度或程序暂停可能参与迟到。现有记录没有网卡硬件时间戳或内核到包时间，不能证明包早已到达而应用未及时读取，也不能排除网络路径延迟。该关联作为定位线索保留。')
h('6 2 服务端证据与统计边界',2)
p('10月13日server.log存在最终逐session汇总。3轮直连与48轮WAGO均找到对应session，服务端唯一接收数与发送端尝试数一致，匹配流重复数及已观察序号缺口均为0。该批回包证据与接收记录一致。')
p('10月14日server2.log结束仍为每秒进度，没有最终逐session汇总。最后累计缺口1856、未对账流62为整批统计，未与本报告48轮一一对应，不能直接记为这些Weidmuller轮次的丢包。发送端CSV中全部尝试最终取得回包，是本批轮次已回包的直接证据。服务器仍运行时，后续正常停止并保存最终汇总【待补】。')
p('server.log的内核IPv4丢弃增量15022为全网卡统计，不能归因到本测试流或交换机。echo没有flood式结束总数对账，不能把服务端“未对账流”直接视为echo失败。')
h('6 3 复核优先顺序',2)
table(['优先级','设备及端口对','先测帧长','复核重点'],[
['1','Weidmuller 13—1','1518B、64B','跨帧异常及连续2包'],
['2','两台 2—1','512B','相同组合两台均迟到'],
['3','WAGO 11—1、16—1','64B、512B','分别约127、124ms迟到'],
['4','Weidmuller 3—1、WAGO 5—1','64B、1518B','单包迟到']],[1.4,5.5,3.1,7.0])
p('各项先进行原接线600秒重复及另一台同配置对照。复核后分别填写“已复现”“本轮未复现”及“原因判断”。Weidmuller 6—1三种帧长及11—1的64B虽无超时，但p99约4.289～5.433ms，若业务对毫秒级时延敏感，可排在超时复核之后补测。')

page();h('7 试验结论与后续安排')
h('7 1 阶段性结论',2)
p('在本次60秒、单流、目标1000pps、50ms超时门槛的条件下，两台交换机16个端口均已覆盖测试，全部尝试最终收到回包，未记录本地发送失败和socket错误。WAGO出现4个迟到包，Weidmuller出现5个，其中Weidmuller 13—1的1518B出现连续2包迟到。因此已取得全口短测回包证据，但异常配置的时限满足情况仍需复核。')
p('异常同时伴随较长发包间隙，测量端暂停是需要检查的候选因素。当前证据无法将事件确定归因到单个端口、交换机、线缆或PC。两台主体时延接近，Weidmuller尾部时延较高；由于测量日期与现场条件记录存在限制，设备性能差异待交替配对复测确认。')
h('7 2 尚待完成的验证',2)
p('完成直连每种帧长600秒长基线，补齐异常配置与另一台设备的同条件复核，再完成1—8、1—16、8—16及生产关键口的600秒测试。对持续异常、两台表现有差异或生产关键组合选择1800秒验证。')
p('补录设备型号编号、实际协商速率、网卡与线缆信息、软件哈希、现场操作及服务端最终汇总；确认业务允许停顿、正式判据与替换对象。真实业务掉站或报警记录如有，应另列用例与结果，不用实验室低负载回显代替生产验证。')
h('7 3 最终结论填写区',2)
p('截至【日期时间待补】，累计完成直连【待补】轮、WAGO【待补】轮、Weidmuller【待补】轮；最长单轮【待补】分钟。异常复核结果为【待补】，已确认原因及证据为【待补】，尚未解释事件为【待补】。')
p('在【实际负载、端口范围、时限及试验环境待补】条件下，设备对业务时限及转发可靠性要求的满足情况为【待补】。结合【生产验证或报警记录待补】，使用或替换意见为【待补】，适用限制为【待补】。')
h('附录说明')
p('附录A、B记录两台各16组接线的三种帧长短测结果，单元格格式为“尝试数 / 超时数”。本批全部超时均为最终收到的迟到包，未回包及本地发送失败均为0。异常值以“＊”标记，具体包号见第5章。逐轮参数、时延及文件名索引保存在out/analysis_20261014/round_summary.csv，9个异常包保存在exception_packets.csv。')

h('附录C 原始证据索引与补充记录')
p('DIRECT目录为直连记录，WAGO与Weidmuller目录为相应交换机记录，BATCH目录为服务端日志。文件前缀日期时间为批次标签；各帧实际启动时间以日志“启动时间”字段为准。报告本版结果对应既有99轮分析快照，进行中的新轮次完成后另行追加。')
p('复核记录建议字段：设备、A口、B口、帧长、实际时长、轮次、完整参数、尝试数、迟到数、未回包数、最长连续超时、p50/p99/最大收到RTT、现场动作、文件编号、复核判断。')
p('方法依据：项目《国产交换机短时测试操作规程》《短时测试判据建议》及SwitchTest.py / switch_core.py实现。正式判据批准人、版本与生效日期为【待补】。')

for dev, label in [('WAGO','A'),('Weidmuller','B')]:
    doc.add_page_break();h(f'附录{label} {dev}全端口短测结果')
    p('每帧长正式发包60秒，timeout 50ms。N为尝试数，L为超时数；＊表示本轮出现超时。')
    rows=[]
    for a in range(1,17):
        b=2 if a==1 else 1
        rs=[next(r for r in rounds if (r['device'],r['a'],r['b'],r['frame'])==(dev,a,b,f)) for f in (64,512,1518)]
        rows.append([f'A{a:02}—B{b:02}']+[f"{r['n']:,} / {r['late']+r['unreturned']+r['send_failed']}{'＊' if r['late'] else ''}" for r in rs]+['需复核' if any(r['late'] for r in rs) else '本轮无超时'])
    table(['端口对','64B N / L','512B N / L','1518B N / L','结果'],rows,[3.0,3.6,3.6,3.6,3.2],9)
    p('每个端口对三种帧长均只完成一次短测。表中结果仅支持已测时长与负载下的观察，未出现超时不证明所有负载或长期运行均无异常。')

doc.core_properties.title='WAGO与Weidmuller交换机转发可靠性对比试验报告'
doc.core_properties.subject='全端口短测结果及长测复核记录'
doc.core_properties.author=''
doc.save(DEST/'交换机转发可靠性对比试验报告_V0.1.docx')
(DEST/'交换机转发可靠性对比试验报告_V0.1.md').write_text('\n'.join(md),encoding='utf-8')
print(DEST/'交换机转发可靠性对比试验报告_V0.1.docx')
