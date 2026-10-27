"""Build a traceable Chinese replacement assessment from recomputed evidence."""
import csv
import json
import statistics
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
DATA = ROOT/'out/replacement_analysis_20261026'
DEST = ROOT/'报告'
STEM = 'WAGO替换Weidmuller综合数据分析与评估_20261026'
E = json.loads((DATA/'echo.json').read_text(encoding='utf-8'))
F = json.loads((DATA/'flood.json').read_text(encoding='utf-8'))
C = json.loads((DATA/'capture_summary.json').read_text(encoding='utf-8'))
S = json.loads((DATA/'summary.json').read_text(encoding='utf-8'))
rows=E['rounds']
BRANDS=('WAGO','Weidmuller')
doc=Document()
md=[]
sec=doc.sections[0]
sec.page_width,sec.page_height=Cm(21),Cm(29.7)
sec.top_margin,sec.bottom_margin,sec.left_margin,sec.right_margin=Cm(1.8),Cm(1.8),Cm(1.8),Cm(1.8)
sec.footer_distance=Cm(.8)
for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
    style=doc.styles[name]
    style.font.name='Microsoft YaHei'
    style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.line_spacing=1.2
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Title'].font.size=Pt(23)
doc.styles['Title'].paragraph_format.space_after=Pt(13)
doc.styles['Heading 1'].font.size=Pt(16)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(9)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(8)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(6)
foot=sec.footer.paragraphs[0]
foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
foot.add_run('2026年10月26日  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
for r in foot.runs:r.font.size=Pt(8)


def p(text,bold=False):
    par=doc.add_paragraph()
    par.add_run(text).bold=bold
    md.extend([text,''])
    return par


def h(text,level=1):
    doc.add_heading(text,level=level)
    md.extend(['#'*(level+1)+' '+text,''])


def page():
    doc.add_page_break()


def table(headers,body,widths=None,size=9):
    t=doc.add_table(rows=1,cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    if widths is None:widths=[17.4/len(headers)]*len(headers)
    for i,w in enumerate(widths):
        t.columns[i].width=Cm(w)
    for vals in [headers,*body]:
        cells=t.rows[0].cells if vals is headers else t.add_row().cells
        for i,(cell,val) in enumerate(zip(cells,vals)):
            cell.width=Cm(widths[i])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=str(val)
            pr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E5EBF1' if vals is headers else 'FFFFFF');pr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge,value in [('top',65),('bottom',65),('left',85),('right',85)]:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),str(value));el.set(qn('w:type'),'dxa');margins.append(el)
            pr.append(margins)
            for par in cell.paragraphs:
                par.paragraph_format.space_after=Pt(0)
                par.paragraph_format.space_before=Pt(0)
                par.paragraph_format.line_spacing=1.08
                par.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 or len(str(val))>20 else WD_ALIGN_PARAGRAPH.CENTER
                for run in par.runs:
                    run.font.size=Pt(size)
                    run.bold=vals is headers
                    run.font.color.rgb=RGBColor(0,0,0)
        trpr=cells[0]._tc.getparent().get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit');trpr.append(no_split)
    headpr=t.rows[0]._tr.get_or_add_trPr()
    repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');headpr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    md.append('| '+' | '.join(headers)+' |')
    md.append('| '+' | '.join('---' for _ in headers)+' |')
    md.extend('| '+' | '.join(str(x).replace('\n',' / ') for x in r)+' |' for r in body)
    md.append('')


def link(label,target):
    par=doc.add_paragraph()
    hyperlink=OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'),par.part.relate_to(target,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r');pr=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'1F4E79');pr.append(col);run.append(pr)
    text=OxmlElement('w:t');text.text=label;run.append(text);hyperlink.append(run);par._p.append(hyperlink)
    md.extend([f'[{label}]({target})',''])


def fmt(n):return f'{n:,}'
def label(a,b):return f'A{a:02}—B{b:02}'
def rr(b,d):return [r for r in rows if r['brand']==b and r['duration_s']==d]


doc.add_paragraph('WAGO替换Weidmüller交换机\n综合数据分析与评估','Title')
md.extend(['# WAGO替换Weidmüller交换机综合数据分析与评估',''])
p('分析日期 2026年10月26日    数据期 2026年10月13日至26日')
p('被测设备：WAGO 852-116；Weidmüller IE-SW-ELB-16-16TX-QL。')
h('1 替换结论')
p('当前建议：将WAGO列为替代候选，完成针对性补测后进入可回退的小范围业务验证；现阶段不建议直接批准全量替换。',True)
p('现有数据支持WAGO具备基本转发能力，并在全口短测尾部时延和高负载异常出现范围上表现较好。但WAGO仍有重复出现的迟到和较重的启动缺失，无法据此确认其在现场业务中与Weidmüller等效。两台测试链路都出现未到达，责任设备尚未定位，原设备也不能作为已证明无异常的基准。')
table(['比较项','WAGO','Weidmüller'],[
    ['全口60秒ECHO','48轮；4包迟到','48轮；5包迟到'],
    ['上述ECHO最终未回包','0','0'],
    ['短测p99轮次中位数','1.837～1.989 ms','2.277～3.380 ms'],
    ['600秒ECHO复测','12轮；4包迟到','10轮；668包迟到'],
    ['共同15组首次100M FLOOD有缺失','3 / 15组','10 / 15组'],
    ['共同15组首次100M FLOOD未到达','204,519 / 7,301,041','101,228 / 7,308,286'],
],[6.3,5.55,5.55])
p('这里的“有缺失组数”和“缺失包数”给出不同排序：WAGO异常较集中，但启动阶段单次缺失较大。长测端口选择、复测次数和测试日期也不一致，因此不能用4与668、3与10这类总数直接证明品牌可靠性高低。')
p('若必须现在作出是否直接替换的决定，结论是暂缓正式全量替换。继续评估WAGO有数据依据；判定WAGO不能替换或已经全面优于原设备，都超出了当前证据。')

page();h('2 数据范围与判定口径')
table(['数据集','实际归档范围','本次核验'],[
    ['ECHO短测','直连3轮；两台各48轮，均60秒','99轮逐包CSV及日志'],
    ['ECHO长测','WAGO 12轮；Weidmüller 10轮，均600秒','22轮含R03复测'],
    ['FLOOD','WAGO 24轮；Weidmüller 17轮','41份最终对账日志'],
    ['双端抓包','14组配对，共28个pcapng','约20.21 GB逐包重新解析'],
    ['服务端批次日志','out/BATCH内5份日志','121个ECHO会话均匹配最终统计'],
],[3.1,7.5,6.8])
p('两台ECHO均覆盖A02—B01至A16—B01以及A01—B02的16组接线，每组64、512、1518B三种请求帧长。FLOOD为1518B单向UDP：WAGO包含上述16组；Weidmüller现存记录为A02—B01至A16—B01，缺少A01—B02方向。1号口已作为B端接收口参与测试，这与补齐反向接线不是同一项覆盖。')
p('ECHO参数一致：单流、目标间隔1 ms、timeout 50 ms、drain 0.5 s、warmup 3 s。CSV包号连续、worker为0、发送时刻单调，逐轮尝试数与日志一致，迟到附件匹配；本地发送失败和socket错误均为0。121个服务端匹配会话的唯一接收数与客户端尝试数相等，重复数和已观察缺口均为0。旧报告“服务端最终汇总未齐”的描述已不适用于当前归档。')
h('采用修订版规则',2)
p('采用《交换机转发可靠性对比试验报告_V0.2_判定规则修订》3.3节：记录不完整或测试中止为INVALID；完整测试中有unreturned或send_failed为FAIL；其余情况下最长连续迟到至少10包为INVALID，1～9包为REVIEW，全部按时回包为PASS。该项目规则用于ECHO轮次判读，不替代现场业务验收标准。')
p('旧Markdown建议稿的0.001%界限与修订版不同，本报告不混用。因连续迟到被判INVALID，不表示数据损坏、事件不存在或硬件无责；原始计数照常保留。FLOOD最终有未到达的轮次记FAIL；零未到达只记该项观测正常，发送端日志缺失时不补判完整PASS。')
h('比较边界',2)
p('低负载回显测得的是两台PC及网络路径的应用层往返时间，回显仅返回测试头，不能等同交换机内部转发时延或双向满帧吞吐。1000 pps对应64、512、1518B请求的估算线速约0.672、4.256、12.304 Mbps，远低于百兆满载。全口依次短测也不等于16口并发满载。')

page();h('3 全口短测时延与回包结果')
table(['设备','轮数','尝试数','PASS','REVIEW','迟到包','未回包'],[
    [b,len(rr(b,60)),fmt(sum(r['n'] for r in rr(b,60))),sum(r['verdict']=='PASS' for r in rr(b,60)),
     sum(r['verdict']=='REVIEW' for r in rr(b,60)),sum(r['late'] for r in rr(b,60)),sum(r['unreturned'] for r in rr(b,60))]
    for b in ('DIRECT',*BRANDS)
],[3.0,1.0,3.1,1.8,2.0,2.0,2.5])
p('短测共91轮PASS、8轮REVIEW。9个迟到包最终全部收到，不能把它们写成9个交换机丢包。两台均取得了全口低负载回包证据。')
table(['帧长B','W p50','原 p50','W p90','原 p90','W p99','原 p99'],[
    [frame,*[f'{S["echo"][b]["60"]["quantiles_by_frame"][str(frame)][key]:.3f}'
        for key in ('p50_ms','p90_ms','p99_ms') for b in BRANDS]] for frame in (64,512,1518)
],[2.4,2.5,2.5,2.5,2.5,2.5,2.5])
p('单位ms。W为WAGO，原为Weidmüller。每个数为同帧长16轮分位数的中位数，不是所有包合并后的分位数；按时样本受50 ms门槛截断。分位数使用排序后(n−1)×p位置的线性插值。')
reductions=[(1-S['echo']['WAGO']['60']['quantiles_by_frame'][str(f)]['p99_ms']/S['echo']['Weidmuller']['60']['quantiles_by_frame'][str(f)]['p99_ms'])*100 for f in (64,512,1518)]
p(f'WAGO的p50低约2%～5%，p99分别低{reductions[0]:.1f}%、{reductions[1]:.1f}%和{reductions[2]:.1f}%，是本批数据中明确的尾部时延优势。不过两台短测分日期完成，每个端口和帧长只有一轮，因此这是观测优势，尚未证明是设备固有差异。')
table(['短测异常','WAGO','Weidmüller'],[
    ['端口与帧长','A02 512B；A05 1518B；A11 64B；A16 512B','A02 512B；A03 64B；A13 64B及1518B'],
    ['最大收到RTT','126.759 ms','117.532 ms'],
    ['最长连续迟到','1包','2包'],
],[3.4,7.0,7.0])
p('WAGO并非每个指标都占优：其短测原始最大RTT高于Weidmüller。8个短测异常轮次均伴随较长发送间隙，提示端点调度可能参与，但不能由相关性直接给PC或交换机定责。')
p('直连仅有三种帧长各60秒，p50约0.596～0.679 ms。未找到600秒直连ECHO或同速率直连FLOOD基线；两端直连可能协商千兆，而经被测交换机为百兆，因此不能把直连与交换机路径的时延差全部算成交换机处理时延。')

page();h('4 十分钟复测与异常是否关闭')
longbody=[]
for b,a in [('WAGO',2),('WAGO',5),('WAGO',11),('WAGO',16),('Weidmuller',2),('Weidmuller',3),('Weidmuller',13)]:
    rs=[r for r in rr(b,600) if r['port_a']==a]
    verdicts=', '.join(f'{name} {sum(r["verdict"]==name for r in rs)}' for name in ('PASS','REVIEW','INVALID') if any(r['verdict']==name for r in rs))
    longbody.append([f'{b} A{a:02}—B01',len(rs),fmt(sum(r['n'] for r in rs)),sum(r['late'] for r in rs),verdicts])
table(['设备及端口对','轮数','尝试数','迟到包','状态'],longbody,[5.0,1.1,3.3,1.5,6.5])
p('WAGO共7,183,291次尝试，10轮PASS、2轮REVIEW；Weidmüller共5,984,737次尝试，9轮PASS、1轮INVALID。全部长测最终回包收到，本地发送失败为0。每台各自重点端口选择不同，长测汇总不能作为同条件配对胜负表。')
h('WAGO A16仍有待解释的时限异常',2)
p('512B R02连续2包迟到，最大53.089 ms，附近发送间隙约51.998 ms；1518B R02也有连续2包迟到，最大51.261 ms，但迟到包附近发送间隔约1.003 ms，本轮最大发送间隙仅17.156 ms。后一事件不能套用“50 ms以上发送暂停”解释，仍需检查收发两端与网络路径。两轮均为REVIEW。')
p('A02、A05、A11三种帧长的600秒复测未出现迟到。可写“本轮未复现”，不能倒推首次异常无效。A16的512B在短测与长测中都出现迟到，因此优先级高于已通过一次长测的其他组合。')
h('Weidmüller A03出现过大批迟到',2)
p('A03—B01、512B R02共668包迟到，最长连续496包，最大收到RTT为2452.649 ms，最大发送间隙为2406.128 ms；存在明显端点暂停线索。依修订规则该轮为INVALID，保留全部事件，不能当作已证明交换机中断2.45秒。')
p('随后同端口同帧长R03运行600秒，598,756次尝试全部按时回包。A02、A13的600秒复测也无迟到。这削弱了“固定端口必然故障”的推断，但仍不能抹去首次事件或证明异常根因已解决。')
p('两台真正共同的600秒端口对只有A02—B01（三种帧长），六轮均无迟到。现有长测因此尚不构成覆盖两台相同重点端口的完整交替对照。')

page();h('5 高负载测试的双重结果')
table(['统计范围','WAGO','Weidmüller'],[
    ['全部FLOOD轮次','24轮：19轮零未到达，5轮FAIL','17轮：6轮零未到达，11轮FAIL'],
    ['其中100 Mbps目标','22轮：18轮零未到达，4轮FAIL','17轮：6轮零未到达，11轮FAIL'],
    ['共同15组首次100M测试','12组零未到达，3组有缺失','5组零未到达，10组有缺失'],
    ['共同组未到达总数','204,519','101,228'],
    ['共同组提交总数','7,301,041','7,308,286'],
    ['共同组汇总未到达比例',f'{204519/7301041*100:.3f}%',f'{101228/7308286*100:.3f}%'],
],[6.2,5.6,5.6])
p('公平比较采用两台共同的A02—B01至A16—B01，在1518B、100 Mbps目标、文件标称60秒条件下，每组按时间选最早一轮。这样不会因WAGO额外做了更多正常复测而扩大表面优势。它仍然不是随机交替实验；缺失比例是端到端应用结果，不是硬件转发丢包率。')
h('WAGO异常范围较小但启动缺失较重',2)
p('WAGO首次100M有缺失的三组为A04、A14、A15；A14与A15首轮的201,273包属于已定位的启动前缀，占共同组204,519包缺失的98.4%。A04剩余3,246包无对应异常轮抓包，原因未定。Weidmüller首次100M有缺失的10组分布更广，应用缺失总量则较低。')
p('所有已归档轮次中，WAGO未到达250,254包，可定位启动前缀246,953包，其余3,301包；Weidmüller未到达101,249包，可定位启动前缀100,480包，其余769包。“其余”包括原因未定及恢复附近缺失，不能直接改名为稳定运行丢包，更不能扣掉启动样本后直接宣布合格。')
h('复测结果支持继续评估但尚未闭环',2)
p('WAGO的A04 R02、A14 R02、A15 R03均为零未到达，因此每组接线至少有一轮100M目标零未到达记录。Weidmüller A05 R02也归零，而A02 R02仍有21包未到达。正常复测和首次异常同时保留，不用最后一次结果覆盖全过程。')
p('WAGO A02的80 Mbps目标轮次无缺失、95 Mbps轮次缺36包、100 Mbps轮次无缺失，单次不同速率记录未形成稳定的负载响应曲线。没有同速率直连FLOOD基线，也没有41轮对应发送端日志，因此无法核实每轮实际线速、本地提交失败和正式发送参数。日志标注100M不是已证明无损百兆吞吐。')

page();h('6 逐端口结果对照')
p('ECHO列为60秒短测三种帧长的迟到数，顺序64 / 512 / 1518B。FLOOD列为100M目标各次归档的未到达数，按时间排列，箭头只表示测试先后。全部ECHO短测最终未回包为0。')
portrows=[]
for a in range(1,17):
    b=2 if a==1 else 1
    rec=[label(a,b)]
    for brand in BRANDS:
        es=[next(r for r in rows if r['brand']==brand and r['duration_s']==60 and r['port_a']==a and r['port_b']==b and r['frame_bytes']==frame)['late'] for frame in (64,512,1518)]
        rec.append(' / '.join(map(str,es)))
    for brand in BRANDS:
        fs=sorted([r for r in F if r['brand']==brand and r['port_a']==a and r['port_b']==b and r['target_mbps']==100],key=lambda r:r['stamp'])
        rec.append(' → '.join(fmt(r['missing']) for r in fs) if fs else '未归档')
    portrows.append(rec)
table(['端口对','W ECHO迟到','原 ECHO迟到','W FLOOD未到达','原 FLOOD未到达'],portrows,[2.8,3.0,3.0,4.3,4.3],size=8.8)
p('WAGO A03现存三份100M日志均标R01，以时间戳和session识别，不误去重为一轮；项目记录中提到的历史78包缺失因原始日志已删除，未进入定量统计，正常记录不能证明该历史事件不存在。')
p('WAGO A14首轮session 0a453a39的原日志名写作A13，沿用项目补充记录中的人工接线纠错，按A14统计，文件名保持原样。A口与B口是电脑角色；异常组合共享B01，不能单靠换A口把责任定位为某一个物理口。')

page();h('7 抓包与应用日志交叉定位')
table(['设备 端口 轮次','应用缺失','启动前缀','其余缺失','应用收到减抓包','抓包丢弃'],[
    [f'{"W" if r["brand"]=="WAGO" else "原"} A{r["port_a"]:02} {r["round"]}',fmt(r['app_missing']),fmt(r['startup_prefix']),fmt(r['other_missing']),fmt(r['app_minus_server_capture']),fmt(r['server_drop'])]
    for r in C
],[3.4,2.6,2.6,2.4,3.4,3.0],size=8.8)
p('14个配对会话的结束总数均与服务端日志分母一致，共同序号的内嵌发送时间戳无不一致。接收应用唯一数减去接收抓包唯一F包数，均等于抓包接口报告丢弃数；这是计数核验，不能恢复每个漏抓包的身份。')
p('新增核验的WAGO A04 R02应用零未到达，但服务端抓包漏记81,978包，是“抓包缺包不等于应用丢包”的直接例子。Weidmüller A15应用缺7包、抓包丢弃61,914包，两者也必须分别保留。存在抓包丢弃，不能据此把真实应用缺失自动清零。')
p('WAGO A15 R01及R02的预热回包延迟达约14～15秒和7.8～8.2秒，之后才进入正常通信；Weidmüller A02 R01也有约7秒预热回包延迟。另有A14类双端均未捕获起始正式包。发送端抓到报文仍只证明抓包点看到提交过程，不证明已被交换机接收。')
p('Weidmüller A05、A08、A09、A13、A15、A16的预热已正常回包，随后仍分别有9、218、84、73、7、77包应用缺失，不能用初始ARP未建立解释全部异常。WAGO A15 R02也有启动前缀之外的19包缺失。')
p('当前程序FLOOD预热按固定时长发送W报文，没有读取并验证及时回包后才开始正式计数。该准备判定缺口会让未就绪状态进入测量，但底层启动延迟及后续缺失的原因仍未查明。候选路径包括发送驱动与网卡、线缆、交换机及参考口、接收驱动与网卡、系统和应用。')

page();h('8 两台交换机的优势与不足')
h('WAGO 852 116',2)
p('优势：16组接线的低负载ECHO覆盖完整，短测与长测全部尝试最终回包；全口短测p50略低，p99中位数低约19%～41%；共同15组首次100M测试中的异常组更少，并且所有16组已有至少一轮100M目标零未到达记录。这些结果支持继续推进替代验证。')
p('不足：启动前缀缺失规模较大，导致首次同范围的汇总未到达比例高于Weidmüller；A16在长测512B和1518B都出现连续2包迟到，不能视为全部重点口复核通过；A04初次3,246包和A02的95M轮次36包等事件尚未定位。价格、寿命、抗干扰和维护成本没有本次实测证据，不作优劣结论。')
h('Weidmüller IE SW ELB 16 16TX QL',2)
p('优势：同样完成全口低负载ECHO，所有尝试最终回包；A02与A13长测正常，A03大批迟到后的同配置R03也正常；共同首次100M样本的未到达总量低于WAGO。保留原设备有利于继续开展同条件对照。')
p('不足：本批短测尾部时延较高，100M目标应用缺失分布更广；A03出现过668包迟到及最长496包连续迟到，A02 FLOOD复测仍缺21包；A01—B02高负载方向未找到归档记录。上述问题都不能直接当作已确认的交换机硬件缺陷。')
h('硬件和工业协议等效性仍需独立核对',2)
p('用户确认的原设备带-QL后缀。公开官方资料检索得到的IE-SW-ELB-16-16TX（无-QL）不能自动用于这台设备的电源、温度、过滤或认证比较，本报告未把无-QL参数当作被测机已确认规格。')
p('WAGO官方产品页确认852-116为16口100Base-TX；官方2026产品表列出12～48 V DC、50×160×120 mm、−40～+70°C和IEEE 802.1p。该表PROFINET栏为“—”，只能说明此表未列CC-A，不能据此推断不能转发PROFINET，也不能把优先级功能等同于系统兼容性证明。')
link('WAGO 852-116官方产品页','https://www.wago.com/gb/industrial-switches/industrial-eco-switch/p/852-116')
link('WAGO官方产品表 第14页 2026年资料','https://lp.wago.com/hubfs/SO_0001-WAGO%20GmbH%20%20Co.%20KG/0001_DE/DE-DE_gated-content/2026_WAGO_Verkaufsaktion_flyer.pdf?hsLang=de-de')
p('替换前需用两台准确版本的资料和现场配置确认供电、安装空间、温度、QoS及所需工业协议；业务类型、报文周期、允许停顿和掉站判据尚未给出。本次UDP测试未覆盖这些应用要求，50 ms为测试门槛，不是已确认的生产允许中断时间。')

page();h('9 最小补测方案与放行条件')
p('建议优先完成以下验证，不需要再次无差别重复全部既有低负载短测。新增测试同时保存发送端与接收端日志、完整参数、设备及网卡信息、实际协商速率、现场动作和程序版本。')
table(['优先级','验证内容','完成后能回答的问题'],[
    ['1','同为100Mbps全双工的直连：64/512/1518B ECHO各600秒；1518B FLOOD在现场所需速率及100M目标下对照','端点链路在相同速率下能否稳定测量'],
    ['2','正式FLOOD前确认当前会话连续及时回包；记录冷启动、换线后和闲置后启动场景','启动未就绪是否仍反复出现'],
    ['3','两台同接线交替测A16的512/1518B ECHO，以及A04/A14/A15的FLOOD；异常项至少3组配对','WAGO待解释事件是否复现，差异是否随设备变化'],
    ['4','Weidmüller A03 512B ECHO与A02/A13/A16 FLOOD配对复核；补A01—B02方向','原设备异常及高负载方向是否闭环'],
    ['5','若复现，逐项更换B01参考口、线缆或接收网卡，再对换收发角色；改善抓包可靠性','问题是否随共同部件或路径变化'],
    ['6','真实PLC及实际协议、实际端口负载下验证启动、稳态及恢复；先做可回退的小范围试用','现场通信周期、掉站和报警要求是否满足'],
],[1.3,9.3,6.8],size=9)
p('端口计数、镜像或独立采集点仅在设备支持时使用；本项目现有程序可导出新的FLOOD缺失序号清单，应与双端抓包按session对账。历史日志未保存完整接收位图，无法补造旧轮次精确缺失包号。')
h('建议的正式放行条件',2)
p('先确定业务允许延迟和中断阈值；有效基线及同条件配对测试中无未解释的未到达、本地失败或越过业务时限的事件；WAGO A16迟到与FLOOD启动异常得到定位或通过有明确条件的复核；准确型号的供电、环境和所需协议功能满足现场要求；小范围业务验证完成且无新增掉站或报警。')
p('上述条件满足后，可在已验证的端口、负载和业务范围内批准替换；若异常仍复现且无法定位，继续暂缓。三组600秒配对属于初步工程复核，不构成长寿命、16口并发容量或长期可靠性证明。')

page();h('10 证据索引与复算说明')
p('原始CSV、日志、pcapng和旧报告未修改。分析脚本按归档文件及会话识别轮次，不按文件中的R01/R02标签盲目合并；同条件首次样本按时间选择，后续复测独立保留。')
table(['输出','用途'],[
    ['echo_rounds.csv','121轮ECHO参数、计数、分位数、判定、服务端匹配及原始文件索引'],
    ['echo_events.csv / echo_failure_segments.csv','681个迟到事件及连续失败段，可回溯包号与实际发送时刻'],
    ['flood_rounds.csv','41轮FLOOD最终对账、原日志和接线纠错说明'],
    ['capture_summary.csv / captures目录','14组双端原始抓包复算、启动缺失、捕获丢弃、预热信息'],
    ['summary.json','分设备、分时长、分速率及共同首次样本汇总'],
    ['source_manifest.json','333份输入与分析代码的SHA-256、大小和修改时间'],
],[6.7,10.7])
p('结构化结果目录：G:\\Codex\\交换机测试\\out\\replacement_analysis_20261026。原始ECHO来自out/DIRECT、out/WAGO和out/Weidmuller；FLOOD来自out/FLOOD；抓包来自WireShark抓包；服务端最终会话来自out/BATCH。out/logs中的本机回环开发测试未当作交换机实测。')
p('规则来源：报告/交换机转发可靠性对比试验报告_V0.2_判定规则修订.docx第3.3节。端口纠错及缺失历史记录来源：报告/FLOOD全部未到达排查汇总_20261026.md。统计共121轮ECHO、41轮FLOOD；它们不能合并为同一丢包率。')
p('ECHO分位数仅针对按时回包，原始最大RTT包含迟到。trimmed_max_ms移除最大ceil(n×0.001)个已收到RTT后取最大，仅供辅助，不替代真实最大值。121轮完整性检查未发现异常，修订规则中的1轮INVALID来自连续迟到而非文件完整性问题。')
p('抓包重新解析全部28个pcapng，检查了会话、结束总数、共同序号的发送时间戳、唯一包数与捕获丢弃；与原先13组已分析结果一致，并新增核验WAGO A04 R02。两台PC抓包时间未同步，不用跨PC绝对时间相减计算单向转发延迟。')
p('复算命令：使用Python 3.11或以上在项目根目录执行 python analyze_replacement.py；只分析归档，不向设备发包。重新生成报告执行 python build_replacement_report.py。原始源文件如发生变化，应重新分析并保留本次快照，而不只重写结论。')
p('规格来源访问日期为2026年10月26日。WAGO资料只用于列出已核实的适配检查项；未采用促销价格作成本结论，也未以无-QL的Weidmüller资料代替被测型号。所有性能优劣判断均限定为当前实验观测。')

assert len(rows)==121 and len(F)==41 and len(C)==14
assert sum(r['late'] for r in rows)==681
assert all(not r['integrity_issues'] for r in rows)
assert all(r['server_summaries'] for r in rows)
assert all(r['app_minus_server_capture']==r['server_drop'] for r in C)
doc.core_properties.title='WAGO替换Weidmüller交换机综合数据分析与评估'
doc.core_properties.subject='2026年10月归档测试数据分析与替换决策'
doc.core_properties.author=''
doc.core_properties.keywords='WAGO 852-116; Weidmuller IE-SW-ELB-16-16TX-QL; ECHO; FLOOD'
doc.save(DEST/(STEM+'.docx'))
(DEST/(STEM+'.md')).write_text('\n'.join(md)+'\n',encoding='utf-8')
with (DATA/'port_comparison.csv').open('w',encoding='utf-8-sig',newline='') as file:
    w=csv.writer(file);w.writerow(['端口对','WAGO ECHO迟到 64/512/1518B','Weidmuller ECHO迟到 64/512/1518B','WAGO 100M FLOOD未到达 按时间','Weidmuller 100M FLOOD未到达 按时间']);w.writerows(portrows)
print(str(DEST/(STEM+'.docx')))
print(str(DEST/(STEM+'.md')))
